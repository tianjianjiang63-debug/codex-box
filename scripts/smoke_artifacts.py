"""Exercise editable file creation, readback, and optional Office rendering."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import tempfile
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]


def smoke(with_render: bool = False) -> dict:
    from docx import Document
    from docx.oxml import OxmlElement
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    from pypdf import PdfReader
    from reportlab.pdfgen.canvas import Canvas

    parent = ROOT / ".work" / "smoke"
    parent.mkdir(parents=True, exist_ok=True)
    folder = Path(tempfile.mkdtemp(prefix="run-", dir=parent))
    word_path = folder / "sample.docx"
    document = Document()
    document.add_heading("云端工具检查", 0)
    document.add_paragraph("中文内容与可编辑公式")
    equation = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "F=ma"
    run.append(text)
    equation.append(run)
    document.add_paragraph()._p.append(equation)
    document.save(word_path)
    if "中文内容" not in "".join(p.text for p in Document(word_path).paragraphs):
        raise AssertionError("DOCX text readback failed.")
    with ZipFile(word_path) as archive:
        if b"<m:oMath>" not in archive.read("word/document.xml"):
            raise AssertionError("Native Word equation is missing.")

    sheet_path = folder / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, "=SUM(A1:A2)"
    workbook.save(sheet_path)
    readback = load_workbook(sheet_path, data_only=False)
    if readback.active["A3"].value != "=SUM(A1:A2)":
        raise AssertionError("XLSX formula readback failed.")
    readback.close()

    slides_path = folder / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "云端幻灯片检查"
    slide.placeholders[1].text = "可编辑文本与讲者备注"
    slide.notes_slide.notes_text_frame.text = "讲者备注检查"
    presentation.save(slides_path)
    loaded = Presentation(slides_path)
    if len(loaded.slides) != 1 or "讲者备注检查" not in loaded.slides[0].notes_slide.notes_text_frame.text:
        raise AssertionError("PPTX slide/notes readback failed.")

    pdf_path = folder / "sample.pdf"
    canvas = Canvas(str(pdf_path))
    canvas.drawString(72, 720, "Codex cloud PDF smoke test")
    canvas.save()
    if "Codex cloud" not in PdfReader(str(pdf_path)).pages[0].extract_text():
        raise AssertionError("PDF text readback failed.")
    rendered = []
    if with_render:
        from render_office import render
        for source in (word_path, sheet_path, slides_path):
            pdf = render(source, folder / ("render-" + source.suffix[1:]), png=True)
            rendered.append(str(pdf))
            if source == sheet_path:
                recalculated = "".join(p.extract_text() or "" for p in PdfReader(str(pdf)).pages)
                if "5" not in recalculated:
                    raise AssertionError("Spreadsheet formula did not visibly recalculate to 5.")
    result = {"passed": True, "folder": str(folder), "structural_checks": ["docx_text_and_native_equation", "xlsx_formula", "pptx_editable_text_and_notes", "pdf_text"], "office_rendered": with_render, "rendered_pdfs": rendered, "visual_review": "still required for user deliverables"}
    (ROOT / ".work" / "artifact-smoke.json").write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    print(json.dumps(smoke(args.render), indent=2, ensure_ascii=False))
